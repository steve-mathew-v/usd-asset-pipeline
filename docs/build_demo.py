"""Generate the demo / how-it-works presentation deck."""

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# formal, light palette
NAVY = "1E293B"
NAVY_DEEP = "0F172A"
ACCENT = "0E7490"
ACCENT_LT = "22D3EE"
SLATE = "475569"
SLATE_LT = "94A3B8"
LIGHT = "F1F5F9"
LINE = "CBD5E1"
WHITE = "FFFFFF"
INK = "1E293B"

HEAD = "Georgia"
BODY = "Calibri"
MONO = "Consolas"

HERE = os.path.dirname(os.path.abspath(__file__))
DEMO = os.path.join(HERE, "demo_images")
GUI = os.path.join(HERE, "images")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_num = [0]


def rgb(h):
    return RGBColor.from_string(h)


def slide(bg=WHITE):
    _num[0] += 1
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(bg)
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, radius=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(s, runs, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        for txt, o in para:
            r = p.add_run()
            r.text = txt
            r.font.name = o.get("font", BODY)
            r.font.size = Pt(o.get("size", 16))
            r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = rgb(o.get("color", INK))
            if "spacing" in o:
                r._r.get_or_add_rPr().set("spc", str(int(o["spacing"] * 100)))
    return tb


def bullets(s, items, x, y, w, h, size=15, color=INK, gap=10, marker=ACCENT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        rm = p.add_run()
        rm.text = "•  "
        rm.font.name = BODY
        rm.font.size = Pt(size)
        rm.font.bold = True
        rm.font.color.rgb = rgb(marker)
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run()
            r1.text = head
            r1.font.name = BODY
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = rgb(color)
            r2 = p.add_run()
            r2.text = rest
            r2.font.name = BODY
            r2.font.size = Pt(size)
            r2.font.color.rgb = rgb(color)
        else:
            r = p.add_run()
            r.text = item
            r.font.name = BODY
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
    return tb


def header(s, kick, title):
    text(s, [[(kick.upper(), {"font": MONO, "size": 12, "bold": True, "color": ACCENT, "spacing": 2})]], 0.9, 0.55, 11, 0.3)
    text(s, [[(title, {"font": HEAD, "size": 29, "bold": True, "color": INK})]], 0.9, 0.9, 11.5, 0.8)


def footer(s):
    text(s, [[("OBJ Asset Pipeline", {"size": 10, "color": SLATE_LT})]], 0.9, 7.05, 6, 0.3)
    text(s, [[(str(_num[0]), {"size": 10, "color": SLATE_LT})]], 12.0, 7.05, 0.8, 0.3, align=PP_ALIGN.RIGHT)


def picture(s, folder, name, sx, sy, sw, sh, cap=None):
    path = os.path.join(folder, name)
    iw, ih = Image.open(path).size
    ratio = iw / ih
    if sw / sh > ratio:
        h = sh
        w = sh * ratio
    else:
        w = sw
        h = sw / ratio
    x = sx + (sw - w) / 2
    y = sy + (sh - h) / 2
    box(s, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=WHITE, line=LINE, line_w=1.0)
    s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if cap:
        text(s, [[(cap, {"size": 12.5, "bold": True, "color": SLATE})]], sx, sy + sh + 0.06, sw, 0.35, align=PP_ALIGN.CENTER)


def node(s, x, y, w, h, title, sub, fill, tcolor, scolor):
    box(s, x, y, w, h, fill=fill, radius=True)
    text(s, [[(title, {"font": BODY, "size": 15, "bold": True, "color": tcolor})]], x, y + 0.35, w, 0.4, align=PP_ALIGN.CENTER)
    text(s, [[(sub, {"size": 11.5, "color": scolor})]], x, y + 0.8, w, 0.5, align=PP_ALIGN.CENTER)


def arrow(s, x, y, w=0.7, h=0.4, color=ACCENT):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(color)
    a.line.fill.background()
    a.shadow.inherit = False


# ---------------------------------------------------------------- 1: title
s = slide(NAVY)
box(s, 0.9, 2.5, 0.14, 2.2, fill=ACCENT_LT)
text(s, [[("DEMONSTRATION", {"font": MONO, "size": 13, "bold": True, "color": ACCENT_LT, "spacing": 3})]], 1.3, 2.55, 11, 0.4)
text(
    s,
    [
        [("OBJ Asset Pipeline", {"font": HEAD, "size": 44, "bold": True, "color": WHITE})],
        [("How It Works", {"font": HEAD, "size": 44, "bold": True, "color": WHITE})],
    ],
    1.3,
    3.05,
    11,
    2.0,
    gap=2,
)
text(s, [[("A walkthrough of the management GUI and the Maya tools", {"size": 18, "color": SLATE_LT, "italic": True})]], 1.3, 5.15, 11, 0.5)
text(s, [[("Steve Mathew Varghese", {"size": 14, "color": WHITE})]], 1.3, 5.9, 11, 0.4)

# ---------------------------------------------------------------- 2: overview
s = slide(WHITE)
header(s, "System overview", "One server, two ways to use it")
text(
    s,
    [[("Artists work in Maya; an administrator manages everything through a desktop GUI. Both talk to a central server, which stores the files on disk and the asset records in MongoDB Atlas.", {"size": 15.5, "color": SLATE})]],
    0.9,
    1.75,
    11.5,
    0.9,
)
ay = 3.2
node(s, 0.9, ay, 3.5, 1.7, "Maya tools", "artists upload &\nimport .obj assets", LIGHT, INK, SLATE)
node(s, 0.9, ay + 2.05, 3.5, 1.5, "Management GUI", "admin: server, users,\nasset browser", LIGHT, INK, SLATE)
node(s, 5.6, ay + 1.0, 3.3, 1.7, "FastAPI Server", "handles uploads,\nlogins, thumbnails", NAVY, WHITE, SLATE_LT)
node(s, 9.9, ay + 1.0, 3.0, 1.7, "MongoDB Atlas", "asset records\n(cloud database)", LIGHT, INK, SLATE)
arrow(s, 4.5, ay + 0.55)
arrow(s, 4.5, ay + 2.6)
arrow(s, 9.0, ay + 1.65)
footer(s)

# ---------------------------------------------------------------- 3: GUI login/server/users
s = slide(WHITE)
header(s, "The management GUI", "Login, server control and user accounts")
picture(s, GUI, "login.png", 0.9, 1.85, 3.7, 3.7, "1. Admin logs in")
picture(s, GUI, "server.png", 4.85, 1.85, 3.7, 3.7, "2. Start / stop the server")
picture(s, GUI, "users.png", 8.8, 1.85, 3.7, 3.7, "3. Create user accounts")
text(
    s,
    [[("The administrator logs in, starts the server with one click, and creates an account for each artist — so every upload is tied to a real user.", {"size": 14, "color": SLATE})]],
    0.9,
    6.0,
    11.5,
    0.8,
)
footer(s)

# ---------------------------------------------------------------- 4: GUI asset browser
s = slide(WHITE)
header(s, "The management GUI", "Browsing uploaded assets")
picture(s, DEMO, "assets_thumbs.png", 5.0, 1.85, 7.5, 4.6)
bullets(
    s,
    [
        ("Thumbnails ", "front & top views, captured automatically on upload"),
        ("Who & where ", "each asset shows its uploader and source tool"),
        ("Ready state ", "mark or unmark an asset as approved"),
        ("Delete ", "remove an asset (admin only)"),
    ],
    0.9,
    2.1,
    3.9,
    4.0,
    size=14.5,
    color=INK,
    gap=14,
)
footer(s)

# ---------------------------------------------------------------- 5: Maya shelf + login
s = slide(WHITE)
header(s, "The Maya tools", "Setup: the OBJPipeline shelf")
picture(s, DEMO, "shelf.png", 0.9, 1.8, 6.0, 3.7, "A custom shelf with one button per action")
picture(s, DEMO, "maya_login.png", 7.1, 1.8, 5.4, 3.7, "First click prompts a login")
text(
    s,
    [[("UPL", {"font": MONO, "size": 13, "bold": True, "color": ACCENT}), (" upload   ", {"size": 13, "color": SLATE}),
      ("IMP", {"font": MONO, "size": 13, "bold": True, "color": ACCENT}), (" import ready   ", {"size": 13, "color": SLATE}),
      ("CHK", {"font": MONO, "size": 13, "bold": True, "color": ACCENT}), (" check ready   ", {"size": 13, "color": SLATE}),
      ("RDY", {"font": MONO, "size": 13, "bold": True, "color": ACCENT}), (" mark ready   ", {"size": 13, "color": SLATE}),
      ("URDY", {"font": MONO, "size": 13, "bold": True, "color": ACCENT}), (" unmark", {"size": 13, "color": SLATE})]],
    0.9,
    5.95,
    11.5,
    0.5,
    align=PP_ALIGN.CENTER,
)
footer(s)

# ---------------------------------------------------------------- 6: Maya upload
s = slide(WHITE)
header(s, "The Maya tools", "Uploading an asset  (UPL)")
picture(s, DEMO, "upload_1.png", 0.9, 1.85, 5.85, 3.6, "1. Pick a .obj file")
picture(s, DEMO, "upload_2.png", 7.0, 1.85, 5.85, 3.6, "2. It uploads and confirms")
text(
    s,
    [[("The artist chooses a .obj file; the tool sends it to the server, records who uploaded it, and quietly captures front and top thumbnails behind the scenes.", {"size": 14, "color": SLATE})]],
    0.9,
    5.95,
    11.5,
    0.8,
)
footer(s)

# ---------------------------------------------------------------- 7: Maya check + mark ready
s = slide(WHITE)
header(s, "The Maya tools", "Sharing: check what's ready  (CHK)  and mark ready  (RDY)")
picture(s, DEMO, "check.png", 0.9, 1.9, 5.85, 3.6, "CHK — lists assets ready to import")
picture(s, DEMO, "Checker.png", 7.0, 1.9, 5.85, 3.6, "RDY — marks an asset as approved")
text(
    s,
    [[("Once an asset is approved, any other artist can see it under CHK and pull it straight into their own scene with IMP.", {"size": 14, "color": SLATE})]],
    0.9,
    6.0,
    11.5,
    0.8,
)
footer(s)

# ---------------------------------------------------------------- 8: end-to-end workflow
s = slide(WHITE)
header(s, "Putting it together", "The end-to-end workflow")
steps = [
    ("1", "Upload", "Artist A sends a .obj\nfrom Maya (UPL)"),
    ("2", "Approve", "Mark it ready when\nfinished (RDY)"),
    ("3", "Discover", "Artist B checks what's\nready (CHK)"),
    ("4", "Import", "Pull it into their\nscene (IMP)"),
]
x = 0.9
for num, title, body in steps:
    box(s, x, 2.6, 2.7, 2.3, fill=LIGHT, radius=True)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.05), Inches(2.85), Inches(0.6), Inches(0.6))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(ACCENT)
    c.line.fill.background()
    c.shadow.inherit = False
    text(s, [[(num, {"font": HEAD, "size": 20, "bold": True, "color": WHITE})]], x + 1.05, 2.87, 0.6, 0.55, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, [[(title, {"font": HEAD, "size": 18, "bold": True, "color": INK})]], x, 3.65, 2.7, 0.5, align=PP_ALIGN.CENTER)
    text(s, [[(body, {"size": 13, "color": SLATE})]], x, 4.15, 2.7, 0.8, align=PP_ALIGN.CENTER)
    if num != "4":
        arrow(s, x + 2.75, 3.6, w=0.5, h=0.35)
    x += 3.05
text(s, [[("No manual file-sharing, no wrong versions — the server always hands over the latest approved asset.", {"size": 15, "italic": True, "color": SLATE})]], 0.9, 5.4, 11.5, 0.6, align=PP_ALIGN.CENTER)
footer(s)

# ---------------------------------------------------------------- 9: summary
s = slide(NAVY)
text(s, [[("SUMMARY", {"font": MONO, "size": 12, "bold": True, "color": ACCENT_LT, "spacing": 2})]], 0.9, 0.8, 11, 0.3)
text(s, [[("What it does", {"font": HEAD, "size": 32, "bold": True, "color": WHITE})]], 0.9, 1.2, 11, 0.8)
bullets(
    s,
    [
        "Artists upload .obj assets from Maya to a shared server",
        "Assets are approved, then imported by anyone — straight into Maya",
        "A desktop GUI manages the server, user accounts and the asset library",
        "Front & top thumbnails are captured automatically on upload",
    ],
    0.9,
    2.3,
    11.5,
    2.6,
    size=17,
    color="E2E8F0",
    marker=ACCENT_LT,
    gap=14,
)
box(s, 0.9, 5.5, 11.5, 0.02, fill=SLATE)
text(
    s,
    [[("Built with  ", {"size": 14, "color": SLATE_LT}),
      ("Python · FastAPI · MongoDB Atlas · PySide6 · Maya Python API · Docker", {"font": MONO, "size": 13, "bold": True, "color": ACCENT_LT})]],
    0.9,
    5.75,
    11.5,
    0.4,
)
text(s, [[("Thank you", {"font": HEAD, "size": 18, "bold": True, "color": WHITE})]], 0.9, 6.4, 11, 0.5)

prs.save(os.path.join(HERE, "Demo_HowItWorks.pptx"))
print("saved Demo_HowItWorks.pptx with", len(prs.slides._sldIdLst), "slides")
