"""Generate the MSc proposal presentation deck (5-8 min version)."""

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# palette
NAVY = "0E1628"
NAVY2 = "1A2540"
TEAL = "14B8A6"
TEAL_LT = "2DD4BF"
AMBER = "F59E0B"
SLATE = "64748B"
SLATE_LT = "94A3B8"
LIGHT = "F1F5F9"
WHITE = "FFFFFF"
INK = "0F172A"

HEAD = "Georgia"
BODY = "Calibri"
MONO = "Consolas"

IMG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "images")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(bg)
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=None, radius=False, shadow=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn("a:effectLst"), {})
        sh = ef.makeelement(
            qn("a:outerShdw"),
            {"blurRad": "90000", "dist": "40000", "dir": "5400000", "rotWithShape": "0"},
        )
        clr = sh.makeelement(qn("a:srgbClr"), {"val": "0F172A"})
        alpha = clr.makeelement(qn("a:alpha"), {"val": "16000"})
        clr.append(alpha)
        sh.append(clr)
        ef.append(sh)
        el.append(ef)
    return shp


def text(s, runs, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        for txt, o in para:
            r = p.add_run()
            r.text = txt
            r.font.name = o.get("font", BODY)
            r.font.size = Pt(o.get("size", 16))
            r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = rgb(o.get("color", INK))
            if "spacing" in o:
                r._r.get_or_add_rPr().set("spc", str(int(o["spacing"] * 100)))
    return tb


def bullets(s, items, x, y, w, h, size=15, color=INK, gap=10, marker=TEAL):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        rm = p.add_run()
        rm.text = "—  "
        rm.font.name = BODY
        rm.font.size = Pt(size)
        rm.font.bold = True
        rm.font.color.rgb = rgb(marker)
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run()
            r1.text = head
            r1.font.name = BODY
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = rgb(color)
            r2 = p.add_run()
            r2.text = rest
            r2.font.name = BODY
            r2.font.size = Pt(size)
            r2.font.color.rgb = rgb(color)
        else:
            r = p.add_run()
            r.text = item
            r.font.name = BODY
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
    return tb


def kicker(s, txt, x, y, color=TEAL):
    text(
        s,
        [[(txt.upper(), {"font": MONO, "size": 12, "bold": True, "color": color, "spacing": 2})]],
        x,
        y,
        8,
        0.3,
    )


def picture(s, name, slot_x, slot_y, slot_w, slot_h, frame=True):
    """Place an image fitted (contained, centered) inside a slot, with a frame."""
    path = os.path.join(IMG_DIR, name)
    iw, ih = Image.open(path).size
    ratio = iw / ih
    if slot_w / slot_h > ratio:
        h = slot_h
        w = slot_h * ratio
    else:
        w = slot_w
        h = slot_w / ratio
    x = slot_x + (slot_w - w) / 2
    y = slot_y + (slot_h - h) / 2
    if frame:
        box(s, x - 0.05, y - 0.05, w + 0.1, h + 0.1, fill=WHITE, shadow=True)
    s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))


# ---------------------------------------------------------------- 1: title
s = slide(NAVY)
box(s, 0.9, 1.05, 0.12, 1.9, fill=TEAL)
text(
    s,
    [[("MSc Computer Animation & Visual Effects  ·  NCCA, Bournemouth", {"font": MONO, "size": 13, "color": TEAL_LT, "spacing": 1})]],
    1.25,
    1.1,
    11,
    0.4,
)
text(
    s,
    [
        [("A Lightweight USD-Based Asset", {"font": HEAD, "size": 38, "bold": True, "color": WHITE})],
        [("Pipeline for Cross-DCC Collaboration", {"font": HEAD, "size": 38, "bold": True, "color": WHITE})],
    ],
    1.25,
    1.6,
    11.4,
    1.8,
    space_after=2,
)
text(
    s,
    [[("Reliable, versioned asset exchange for small studios and student teams", {"size": 18, "color": SLATE_LT, "italic": True})]],
    1.25,
    3.45,
    11,
    0.5,
)
box(s, 1.25, 4.55, 10.8, 0.02, fill=NAVY2)
text(
    s,
    [
        [("Project Proposal", {"size": 14, "bold": True, "color": WHITE})],
        [("Steve Mathew Varghese  ·  Supervisor: [TBC]  ·  [Date]", {"size": 13, "color": SLATE_LT})],
    ],
    1.25,
    4.8,
    11,
    0.9,
    space_after=3,
)

# ---------------------------------------------------------------- 2: problem
s = slide(WHITE)
kicker(s, "The problem", 0.9, 0.7)
text(s, [[("Small teams have no real asset pipeline", {"font": HEAD, "size": 31, "bold": True, "color": INK})]], 0.9, 1.05, 11.5, 0.8)
cards = [
    ("Big studios", "Heavyweight systems (ShotGrid, Perforce, proprietary tools). Powerful, but costly and assume dedicated pipeline staff.", SLATE),
    ("Small teams", "Fall back to manual file-sharing: cloud drives, naming conventions, messaging. No versioning, no “approved” state.", AMBER),
    ("The cost", "Wrong versions used, missing textures, overwritten work, time lost chasing the latest file — the errors a pipeline prevents.", TEAL),
]
cx = 0.9
for title_t, body, accent in cards:
    box(s, cx, 2.1, 3.74, 3.6, fill=LIGHT, radius=True, shadow=True)
    box(s, cx + 0.4, 2.5, 0.5, 0.5, fill=accent, radius=True)
    text(s, [[(title_t, {"font": HEAD, "size": 19, "bold": True, "color": INK})]], cx + 0.4, 3.2, 3.0, 0.5)
    text(s, [[(body, {"size": 14.5, "color": SLATE})]], cx + 0.4, 3.75, 2.95, 1.8)
    cx += 3.95

# ---------------------------------------------------------------- 3: prototype + screenshots
s = slide(WHITE)
kicker(s, "Starting point", 0.9, 0.6, color=AMBER)
text(s, [[("I built a working prototype", {"font": HEAD, "size": 31, "bold": True, "color": INK})]], 0.9, 0.95, 11.5, 0.7)

shots = [("login.png", "Login & accounts"), ("users.png", "Multi-user management"), ("server.png", "Server control")]
slot_w, gap = 3.58, 0.4
x = 0.9
for name, caption in shots:
    picture(s, name, x, 1.95, slot_w, 3.5)
    text(s, [[(caption, {"size": 12.5, "bold": True, "color": SLATE})]], x, 5.55, slot_w, 0.35, align=PP_ALIGN.CENTER)
    x += slot_w + gap

box(s, 0.9, 6.15, 11.53, 0.9, fill=NAVY, radius=True)
text(
    s,
    [[("FastAPI + MongoDB server, PySide6 GUI, Maya shelf tools — it works, but ", {"size": 14.5, "color": "E2E8F0"}),
      (".obj and the lack of versioning are exactly its limits.", {"size": 14.5, "bold": True, "color": TEAL_LT})]],
    1.25,
    6.15,
    10.9,
    0.9,
    anchor=MSO_ANCHOR.MIDDLE,
)

# ---------------------------------------------------------------- 4: research question + why USD
s = slide(NAVY)
kicker(s, "The idea", 0.9, 0.65, color=TEAL_LT)
text(
    s,
    [
        [("Can a USD-based system give small teams ", {"font": HEAD, "size": 26, "bold": True, "color": WHITE}),
         ("reliable,", {"font": HEAD, "size": 26, "bold": True, "color": TEAL_LT})],
        [("versioned, cross-DCC", {"font": HEAD, "size": 26, "bold": True, "color": TEAL_LT}),
         (" asset exchange — and beat manual sharing?", {"font": HEAD, "size": 26, "bold": True, "color": WHITE})],
    ],
    0.9,
    1.05,
    11.5,
    1.4,
    space_after=3,
)
box(s, 0.9, 2.75, 5.6, 3.9, fill=NAVY2, radius=True)
text(s, [[(".obj  —  the prototype", {"font": MONO, "size": 16, "bold": True, "color": SLATE_LT})]], 1.3, 3.1, 5, 0.4)
bullets(s, ["Geometry only", "No materials or textures", "No hierarchy or variants", "Re-upload overwrites history"], 1.3, 3.7, 4.9, 2.6, size=15, color="E2E8F0", marker=SLATE_LT, gap=11)
box(s, 6.85, 2.75, 5.6, 3.9, fill=NAVY2, radius=True)
text(s, [[("OpenUSD  —  the proposal", {"font": MONO, "size": 16, "bold": True, "color": TEAL_LT})]], 7.25, 3.1, 5, 0.4)
bullets(
    s,
    [
        ("Versioning ", "via USD composition / layers"),
        ("Cross-DCC ", "native in Maya, Houdini, Unreal"),
        ("Rich data ", "hierarchy + materials travel along"),
        ("Industry standard ", "(Pixar, NVIDIA Omniverse)"),
    ],
    7.25,
    3.7,
    5.0,
    2.6,
    size=15,
    color="E2E8F0",
    marker=TEAL_LT,
    gap=11,
)

# ---------------------------------------------------------------- 5: objectives + methods
s = slide(WHITE)
kicker(s, "Objectives & implementation methods", 0.9, 0.7)
text(s, [[("What I'll build, and how", {"font": HEAD, "size": 31, "bold": True, "color": INK})]], 0.9, 1.05, 11.5, 0.8)

box(s, 0.9, 2.1, 5.6, 4.6, fill=LIGHT, radius=True, shadow=True)
text(s, [[("OBJECTIVES", {"font": MONO, "size": 12, "bold": True, "color": TEAL, "spacing": 2})]], 1.3, 2.45, 5, 0.3)
bullets(
    s,
    [
        "Review existing pipelines (Kitsu, Prism, USD)",
        "Versioned server with an “approved” state",
        "USD publish/consume in Maya + Houdini",
        "Consume “latest approved” or a pinned version",
        "Evaluate with artists (study + SUS survey)",
    ],
    1.3,
    2.95,
    4.9,
    3.5,
    size=14.5,
    color=INK,
    gap=12,
)

box(s, 6.85, 2.1, 5.6, 4.6, fill=NAVY, radius=True)
text(s, [[("METHODS", {"font": MONO, "size": 12, "bold": True, "color": TEAL_LT, "spacing": 2})]], 7.25, 2.45, 5, 0.3)
bullets(
    s,
    [
        ("Python + FastAPI ", "server, USD via pxr bindings"),
        ("PySide6 ", "clients inside Maya & Houdini"),
        ("Docker + CI ", "reproducible deploy & tests"),
        ("Design-and-build ", "+ empirical user evaluation"),
    ],
    7.25,
    2.95,
    5.0,
    2.6,
    size=14.5,
    color="E2E8F0",
    marker=TEAL_LT,
    gap=12,
)
text(s, [[("DCC clients  →  FastAPI server  →  versioned USD storage", {"font": MONO, "size": 13, "bold": True, "color": TEAL_LT})]], 7.25, 6.0, 5.0, 0.4)

# ---------------------------------------------------------------- 6: key challenges
s = slide(WHITE)
kicker(s, "Key challenges", 0.9, 0.7, color=AMBER)
text(s, [[("What could be hard — and the plan", {"font": HEAD, "size": 31, "bold": True, "color": INK})]], 0.9, 1.05, 11.5, 0.8)
challenges = [
    ("USD learning curve", "Powerful but complex. Mitigation: use prebuilt usd-core; a focused week-1 spike."),
    ("Version resolution", "A full custom USD resolver is large. Fallback: resolve versions server-side."),
    ("Persistent file storage", "Cloud disks are ephemeral. Move assets to object storage (S3 / GridFS)."),
    ("Recruiting test artists", "Studies need participants. Recruit from the NCCA cohort; keep tasks short."),
]
positions = [(0.9, 2.1), (6.85, 2.1), (0.9, 4.5), (6.85, 4.5)]
for (px, py), (head, body) in zip(positions, challenges):
    box(s, px, py, 5.6, 2.2, fill=LIGHT, radius=True, shadow=True)
    box(s, px, py, 0.12, 2.2, fill=AMBER)
    text(s, [[(head, {"font": HEAD, "size": 18, "bold": True, "color": INK})]], px + 0.4, py + 0.3, 5.0, 0.5)
    text(s, [[(body, {"size": 14, "color": SLATE})]], px + 0.4, py + 0.95, 5.0, 1.1)

# ---------------------------------------------------------------- 7: references
s = slide(NAVY)
kicker(s, "References & supporting material", 0.9, 0.7, color=TEAL_LT)
text(s, [[("References", {"font": HEAD, "size": 31, "bold": True, "color": WHITE})]], 0.9, 1.05, 11.5, 0.8)
bullets(
    s,
    [
        ("OpenUSD ", "— Pixar / Alliance for OpenUSD (AOUSD) documentation & specification"),
        ("NVIDIA Omniverse ", "— OpenUSD pipeline resources and tooling"),
        ("Autodesk Maya USD / SideFX Houdini Solaris ", "— native USD authoring docs"),
        ("Kitsu (CGWire) & Prism Pipeline ", "— open-source asset/production tools (comparison)"),
        ("Brooke, J. (1996) ", "— SUS: A 'quick and dirty' usability scale (evaluation method)"),
    ],
    0.9,
    2.2,
    11.6,
    3.6,
    size=15.5,
    color="E2E8F0",
    marker=TEAL_LT,
    gap=14,
)
text(s, [[("Thank you — questions welcome.", {"size": 16, "italic": True, "color": SLATE_LT})]], 0.9, 6.5, 11, 0.5)

prs.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "MSc_Proposal.pptx"))
print("saved MSc_Proposal.pptx with", len(prs.slides._sldIdLst), "slides")
